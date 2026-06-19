"""گزارش‌ها و خروجی Excel / PDF."""

import io
from datetime import datetime

from flask import (
    Blueprint, render_template, request, send_file, url_for
)

from ..auth import login_required
from ..db import get_db
from ..jalali import parse_date_input, gregorian_iso_to_jalali, to_persian_digits
from ..utils import format_money

bp = Blueprint("reports", __name__, url_prefix="/reports")

REPORT_TYPES = {
    "projects": "گزارش پروژه‌ها",
    "shop_income": "درآمد مغازه",
    "workers_share": "سهم نیروها",
    "worker_debt": "بدهی به نیروها",
    "unsettled": "پروژه‌های تسویه‌نشده",
}


def _date_range():
    """بازه تاریخ را از فرم می‌خواند و (from_iso, to_iso, from_j, to_j) می‌دهد."""
    from_j, from_iso = parse_date_input(request.values.get("from_date"))
    to_j, to_iso = parse_date_input(request.values.get("to_date"))
    return from_iso, to_iso, from_j, to_j


def _build_report(rtype, from_iso, to_iso):
    """داده گزارش را تولید می‌کند: (سرستون‌ها، ردیف‌ها، خلاصه)."""
    db = get_db()

    def date_clause(field="project_date"):
        clauses, params = [], []
        if from_iso:
            clauses.append("%s >= ?" % field)
            params.append(from_iso)
        if to_iso:
            clauses.append("%s <= ?" % field)
            params.append(to_iso)
        return (" WHERE " + " AND ".join(clauses)) if clauses else "", params

    if rtype == "projects":
        where, params = date_clause()
        rows = db.execute(
            "SELECT p.*, "
            " (SELECT COALESCE(SUM(share_amount),0) FROM project_workers WHERE project_id=p.id) AS workers_share "
            "FROM projects p" + where + " ORDER BY p.project_date, p.id",
            params,
        ).fetchall()
        headers = ["نام پروژه", "مشتری", "تاریخ", "وضعیت", "مبلغ کل", "سهم نیروها", "سهم مغازه"]
        data = []
        sum_labor = sum_ws = 0
        from ..utils import status_label
        for r in rows:
            shop = r["labor_amount"] - r["workers_share"]
            sum_labor += r["labor_amount"]
            sum_ws += r["workers_share"]
            data.append([
                r["name"], r["customer_name"] or "-",
                r["project_date_jalali"] or "-", status_label(r["status"]),
                r["labor_amount"], r["workers_share"], shop,
            ])
        summary = [("تعداد پروژه", len(data)), ("مجموع دستمزد", sum_labor),
                   ("مجموع سهم نیروها", sum_ws), ("مجموع سهم مغازه", sum_labor - sum_ws)]
        return headers, data, summary, [4, 5, 6]

    if rtype == "shop_income":
        where, params = date_clause()
        row = db.execute(
            "SELECT COALESCE(SUM(labor_amount),0) AS labor FROM projects p" + where, params
        ).fetchone()
        where2, params2 = date_clause("p.project_date")
        ws = db.execute(
            "SELECT COALESCE(SUM(pw.share_amount),0) AS s FROM project_workers pw"
            " JOIN projects p ON p.id = pw.project_id" + where2, params2
        ).fetchone()["s"]
        headers = ["شرح", "مبلغ"]
        data = [["مجموع دستمزد پروژه‌ها", row["labor"]],
                ["مجموع سهم نیروها", ws],
                ["درآمد (سهم) مغازه", row["labor"] - ws]]
        summary = [("درآمد مغازه در بازه", row["labor"] - ws)]
        return headers, data, summary, [1]

    if rtype == "workers_share":
        where, params = date_clause("p.project_date")
        rows = db.execute(
            "SELECT w.name, COUNT(pw.id) AS cnt, COALESCE(SUM(pw.share_amount),0) AS share "
            "FROM project_workers pw JOIN projects p ON p.id = pw.project_id "
            "JOIN workers w ON w.id = pw.worker_id" + where +
            " GROUP BY w.id ORDER BY share DESC",
            params,
        ).fetchall()
        headers = ["نیرو", "تعداد پروژه", "مجموع سهم"]
        data = [[r["name"], r["cnt"], r["share"]] for r in rows]
        summary = [("مجموع سهم نیروها", sum(r["share"] for r in rows))]
        return headers, data, summary, [2]

    if rtype == "worker_debt":
        rows = db.execute(
            "SELECT w.name, "
            " COALESCE((SELECT SUM(share_amount) FROM project_workers WHERE worker_id=w.id),0) AS share, "
            " COALESCE((SELECT SUM(amount) FROM payments WHERE worker_id=w.id),0) AS paid "
            "FROM workers w ORDER BY w.name COLLATE NOCASE"
        ).fetchall()
        headers = ["نیرو", "مجموع سهم", "پرداخت‌شده", "مانده طلب"]
        data = []
        total_debt = 0
        for r in rows:
            bal = r["share"] - r["paid"]
            total_debt += bal if bal > 0 else 0
            data.append([r["name"], r["share"], r["paid"], bal])
        summary = [("مجموع بدهی به نیروها", total_debt)]
        return headers, data, summary, [1, 2, 3]

    # unsettled
    rows = db.execute(
        "SELECT p.*, "
        " (SELECT COALESCE(SUM(share_amount),0) FROM project_workers WHERE project_id=p.id) AS workers_share, "
        " (SELECT COALESCE(SUM(amount),0) FROM payments WHERE project_id=p.id) AS paid "
        "FROM projects p WHERE p.status != 'settled' ORDER BY p.id DESC"
    ).fetchall()
    headers = ["نام پروژه", "مشتری", "تاریخ", "مبلغ کل", "سهم نیروها", "پرداخت‌شده", "مانده"]
    data = []
    from ..utils import status_label  # noqa
    for r in rows:
        data.append([
            r["name"], r["customer_name"] or "-", r["project_date_jalali"] or "-",
            r["labor_amount"], r["workers_share"], r["paid"],
            r["workers_share"] - r["paid"],
        ])
    summary = [("تعداد پروژه‌های تسویه‌نشده", len(data))]
    return headers, data, summary, [3, 4, 5, 6]


@bp.route("/")
@login_required
def index():
    rtype = request.args.get("type") or "projects"
    if rtype not in REPORT_TYPES:
        rtype = "projects"
    from_iso, to_iso, from_j, to_j = _date_range()
    headers, data, summary, money_cols = _build_report(rtype, from_iso, to_iso)
    return render_template(
        "reports/index.html",
        report_types=REPORT_TYPES, rtype=rtype,
        headers=headers, data=data, summary=summary, money_cols=money_cols,
        from_j=request.values.get("from_date", ""),
        to_j=request.values.get("to_date", ""),
    )


@bp.route("/print")
@login_required
def print_view():
    """نمای چاپ‌پسند برای ذخیره به PDF از طریق مرورگر."""
    rtype = request.args.get("type") or "projects"
    if rtype not in REPORT_TYPES:
        rtype = "projects"
    from_iso, to_iso, from_j, to_j = _date_range()
    headers, data, summary, money_cols = _build_report(rtype, from_iso, to_iso)
    return render_template(
        "reports/print.html",
        title=REPORT_TYPES[rtype], headers=headers, data=data,
        summary=summary, money_cols=money_cols,
        from_j=request.values.get("from_date", ""),
        to_j=request.values.get("to_date", ""),
        now=gregorian_iso_to_jalali(datetime.now().date().isoformat()),
    )


@bp.route("/export.xlsx")
@login_required
def export_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    rtype = request.args.get("type") or "projects"
    if rtype not in REPORT_TYPES:
        rtype = "projects"
    from_iso, to_iso, _, _ = _date_range()
    headers, data, summary, money_cols = _build_report(rtype, from_iso, to_iso)

    wb = Workbook()
    ws = wb.active
    ws.title = "گزارش"
    ws.sheet_view.rightToLeft = True

    title = REPORT_TYPES[rtype]
    ws.append([title])
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)

    header_fill = PatternFill("solid", fgColor="2F6FED")
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=2, column=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center")

    for row in data:
        out = []
        for idx, val in enumerate(row):
            out.append(val)
        ws.append(out)

    # خلاصه
    ws.append([])
    for label, value in summary:
        ws.append([label, value])

    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max(length + 4, 12), 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = "report_%s_%s.xlsx" % (rtype, datetime.now().strftime("%Y%m%d_%H%M"))
    return send_file(
        buf, as_attachment=True, download_name=fname,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
